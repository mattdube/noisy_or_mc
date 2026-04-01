"""
noisy_or_copula.py
==================
Standalone Gaussian copula version of the Noisy-OR Bayesian Risk Model.

This class accounts for correlations between risk sources using a Gaussian 
copula and redundancy-aware blending. 

Dependencies:
    pip install numpy pandas scipy openpyxl
"""

from __future__ import annotations

import json
import math
import warnings
from datetime import date
from pathlib import Path
from typing import Optional, Union

import numpy as np
import pandas as pd

# Polars is optional
try:
    import polars as pl
    _POLARS_AVAILABLE = True
except ImportError:
    _POLARS_AVAILABLE = False

# ---------------------------------------------------------------------------
# Configuration Constants
# ---------------------------------------------------------------------------
_BATCH_POOL_SIZE = 50_000
_BATCH_CHUNK_BYTES: int = 512 * 1024 * 1024  # 512 MB default memory ceiling


# ---------------------------------------------------------------------------
# Internal Helpers
# ---------------------------------------------------------------------------
def _certainty_label(cert: float) -> str:
    if cert >= 0.80: return "High"
    if cert >= 0.50: return "Medium"
    return "Low"


def _run_batch_vectorized(
    pool: np.ndarray,
    source_means: np.ndarray,
    active_matrix: np.ndarray,
    n_samples: int,
    rng: np.random.Generator,
) -> list[dict]:
    """Score records in a vectorized NumPy pass (baseline Noisy-OR)."""
    n_records, n_sources = active_matrix.shape
    idx = rng.integers(0, pool.shape[1], size=n_samples)
    pool_slice = pool[:, idx]

    p3d = pool_slice[np.newaxis, :, :] * active_matrix[:, :, np.newaxis]
    combined = 1.0 - np.prod(1.0 - p3d, axis=1)

    mean_risk = np.mean(combined, axis=1)
    median_risk = np.median(combined, axis=1)
    p5 = np.percentile(combined, 5, axis=1)
    p95 = np.percentile(combined, 95, axis=1)
    certainty = np.clip(1.0 - (p95 - p5), 0.0, None)

    active_mus = source_means[np.newaxis, :] * active_matrix.astype(float)
    inactive = (active_matrix == 0)
    prod_terms = np.where(inactive, 1.0, 1.0 - active_mus)
    prod_all = np.prod(prod_terms, axis=1)
    overall_risk = 1.0 - prod_all
    denom = np.where(active_mus > 0.9999, np.inf, 1.0 - active_mus)
    prod_without = prod_all[:, np.newaxis] / denom
    contribs_mat = overall_risk[:, np.newaxis] - (1.0 - prod_without)
    contribs_mat = np.where(active_matrix, contribs_mat, 0.0)

    results = []
    for r in range(n_records):
        contribs = {
            int(i): round(float(contribs_mat[r, i]), 4)
            for i in range(n_sources) if active_matrix[r, i]
        }
        results.append({
            "mean_risk": round(float(mean_risk[r]), 4),
            "median_risk": round(float(median_risk[r]), 4),
            "p5": round(float(p5[r]), 4),
            "p95": round(float(p95[r]), 4),
            "certainty": round(float(certainty[r]), 4),
            "certainty_label": _certainty_label(float(certainty[r])),
            "source_means": [round(float(m), 4) for m in source_means],
            "marginal_contribs": contribs,
        })
    return results


def _build_copula_pool(
    mu_values: np.ndarray,
    kappa_values: np.ndarray,
    corr_matrix: np.ndarray,
    pool_size: int,
    rng: np.random.Generator,
) -> np.ndarray:
    """Draw correlated Beta samples using a Gaussian copula."""
    from scipy.stats import norm, beta as beta_dist
    n_sources = len(mu_values)

    try:
        L = np.linalg.cholesky(corr_matrix)
    except np.linalg.LinAlgError:
        eigvals, eigvecs = np.linalg.eigh(corr_matrix)
        eigvals = np.clip(eigvals, 1e-8, None)
        L = np.linalg.cholesky(eigvecs @ np.diag(eigvals) @ eigvecs.T)

    Z = rng.standard_normal((n_sources, pool_size))
    Z_corr = L @ Z
    U = norm.cdf(Z_corr)

    pool = np.empty((n_sources, pool_size), dtype=np.float32)
    for i, (mu, kappa) in enumerate(zip(mu_values, kappa_values)):
        alpha = max(float(mu) * float(kappa), 0.001)
        b = max((1.0 - float(mu)) * float(kappa), 0.001)
        u_clipped = np.clip(U[i], 1e-7, 1.0 - 1e-7)
        pool[i] = beta_dist.ppf(u_clipped, alpha, b).astype(np.float32)
    return pool


# ---------------------------------------------------------------------------
# Main Class
# ---------------------------------------------------------------------------
class NoisyORModelCopula:
    """Noisy-OR Bayesian Risk Model with Gaussian copula correlations."""

    def __init__(self, config_path: Union[str, Path] = "noisy_or_config.json"):
        config_path = Path(config_path)
        if not config_path.exists():
            raise FileNotFoundError(f"Config not found: {config_path}")

        self._config_path = config_path
        self._load_from_config(config_path)
        self._build_batch_pool()

    def _load_from_config(self, config_path: Path) -> None:
        with open(config_path) as f:
            cfg = json.load(f)

        self._sources: list[dict] = cfg["sources"]
        self._n_samples_default: int = int(cfg.get("n_monte_carlo_samples", 20_000))
        self._version: str = cfg.get("version", "1.0.0")
        self._version_date: str = cfg.get("version_date", str(date.today()))

        self._names = [s["name"] for s in self._sources]
        self._mu_values = np.array([s["mu"] for s in self._sources], dtype=float)
        self._kappa_values = np.array([s["kappa"] for s in self._sources], dtype=float)
        self._n_sources = len(self._sources)

        raw_bands = cfg.get("risk_bands", [])
        self._risk_bands: list[tuple[float, str]] = sorted(
            [(float(b["min_score"]), str(b["label"])) for b in raw_bands],
            reverse=True,
        )

        self._corr_matrix, self._corr_pairs = self._parse_correlations(
            cfg.get("correlations", []), self._names
        )

    @staticmethod
    def _parse_correlations(correlations: list[dict], source_names: list[str]):
        n = len(source_names)
        idx = {name: i for i, name in enumerate(source_names)}
        matrix = np.eye(n, dtype=float)
        pairs: list[dict] = []

        for entry in correlations:
            a, b, rho = entry["source_a"], entry["source_b"], float(entry["rho"])
            if a not in idx or b not in idx:
                raise ValueError(f"Unknown source in correlation: {a} or {b}")
            if not (-1.0 < rho < 1.0):
                raise ValueError(f"rho={rho} must be between -1 and 1.")
            i, j = idx[a], idx[b]
            matrix[i, j] = matrix[j, i] = rho
            pairs.append({"source_a": a, "source_b": b, "rho": rho})

        eigvals = np.linalg.eigvalsh(matrix)
        if eigvals.min() < -1e-6:
            raise ValueError(f"Correlation matrix not PSD. Min eigenvalue: {eigvals.min():.4f}")
        return matrix, pairs

    def reload(self) -> None:
        self._load_from_config(self._config_path)
        self._build_batch_pool()
        print(f"Model reloaded: {self!r}")

    # --- Properties (Fixed Missing Attributes) ---
    @property
    def source_names(self) -> list[str]: return self._names

    @property
    def n_sources(self) -> int: return self._n_sources

    @property
    def version(self) -> str: return self._version

    @property
    def version_date(self) -> str: return self._version_date

    @property
    def risk_bands(self) -> list[tuple[float, str]]: return list(self._risk_bands)

    @property
    def correlation_matrix(self) -> np.ndarray: return self._corr_matrix.copy()

    @property
    def correlated_pairs(self) -> list[dict]: return list(self._corr_pairs)

    # --- Methods ---
    def predict(self, active_flags: np.ndarray, n_samples: Optional[int] = None,
                seed: Optional[int] = None, label: Optional[str] = None) -> dict:
        from scipy.stats import norm as _norm, beta as _beta_dist
        active_flags = np.asarray(active_flags, dtype=int)
        self._validate_flags(active_flags)

        n = n_samples or self._n_samples_default
        rng = np.random.default_rng(seed)

        Z = rng.standard_normal((self._n_sources, n))
        try:
            L = np.linalg.cholesky(self._corr_matrix)
        except np.linalg.LinAlgError:
            eigvals, eigvecs = np.linalg.eigh(self._corr_matrix)
            eigvals = np.clip(eigvals, 1e-8, None)
            L = np.linalg.cholesky(eigvecs @ np.diag(eigvals) @ eigvecs.T)

        Z_corr = L @ Z
        U = np.clip(_norm.cdf(Z_corr), 1e-7, 1.0 - 1e-7)

        all_probs = []
        for i, (mu, kappa) in enumerate(zip(self._mu_values, self._kappa_values)):
            alpha = max(float(mu) * float(kappa), 0.001)
            beta = max((1.0 - float(mu)) * float(kappa), 0.001)
            all_probs.append(_beta_dist.ppf(U[i], alpha, beta))

        active_indices = [i for i, a in enumerate(active_flags) if a]
        active_probs = [all_probs[i] for i in active_indices]

        if not active_probs:
            combined = np.zeros(n)
        elif len(active_probs) == 1:
            combined = active_probs[0].copy()
        else:
            nor = 1.0 - np.prod([1.0 - p for p in active_probs], axis=0)
            mx = np.max(active_probs, axis=0)
            rho_eff = self._rho_eff_for_active(active_indices)
            combined = (1.0 - rho_eff) * nor + rho_eff * mx

        source_means = [float(np.mean(p)) for p in all_probs]
        p5, p95 = np.percentile(combined, [5, 95])
        sim = {
            "mean_risk": round(float(np.mean(combined)), 4),
            "median_risk": round(float(np.median(combined)), 4),
            "p5": round(float(p5), 4), "p95": round(float(p95), 4),
            "certainty": round(max(0.0, 1.0 - (p95 - p5)), 4),
            "certainty_label": _certainty_label(max(0.0, 1.0 - (p95 - p5))),
            "source_means": [round(m, 4) for m in source_means],
            "marginal_contribs": self._blended_marginals(active_indices, source_means),
        }
        return self._format_single(sim, active_flags, label)

    def predict_batch(self, active_matrix: np.ndarray, n_samples: Optional[int] = None,
                      seed: Optional[int] = None, labels: Optional[list[str]] = None,
                      backend: str = "pandas") -> Union[pd.DataFrame, pl.DataFrame]:
        active_matrix = np.atleast_2d(np.asarray(active_matrix, dtype=int))
        n = min(n_samples or self._n_samples_default, _BATCH_POOL_SIZE)
        self._validate_matrix(active_matrix)
        rng = np.random.default_rng(seed)

        keys = [tuple(row.tolist()) for row in active_matrix]
        unique_keys = list(dict.fromkeys(keys))
        unique_matrix = np.array(unique_keys, dtype=int)

        unique_sims = _run_batch_vectorized(self._batch_pool, self._batch_source_means, unique_matrix, n, rng)

        idx = rng.integers(0, _BATCH_POOL_SIZE, size=n)
        pool_slice = self._batch_pool[:, idx].astype(float)

        blended_sims = []
        for row_flags, sim in zip(unique_matrix, unique_sims):
            active_idx = [i for i, a in enumerate(row_flags) if a]
            if len(active_idx) < 2 or (rho_eff := self._rho_eff_for_active(active_idx)) < 1e-6:
                blended_sims.append(sim)
                continue

            active_p = pool_slice[active_idx, :]
            combined = (1.0 - rho_eff) * (1.0 - np.prod(1.0 - active_p, axis=0)) + rho_eff * np.max(active_p, axis=0)
            p5, p95 = np.percentile(combined, [5, 95])
            cert = max(0.0, 1.0 - (p95 - p5))

            blended_sims.append({
                **sim, "mean_risk": round(float(np.mean(combined)), 4),
                "median_risk": round(float(np.median(combined)), 4),
                "p5": round(float(p5), 4), "p95": round(float(p95), 4),
                "certainty": round(float(cert), 4), "certainty_label": _certainty_label(cert),
                "marginal_contribs": self._blended_marginals(active_idx, sim["source_means"]),
            })

        sim_lookup = {k: blended_sims[i] for i, k in enumerate(unique_keys)}
        labels = labels or [f"Record_{i+1}" for i in range(len(active_matrix))]
        rows = [self._format_row(sim_lookup[keys[i]], active_matrix[i], labels[i]) for i in range(len(keys))]

        return self._to_polars(rows) if backend == "polars" else pd.DataFrame(rows)

    def export_batch(self, active_matrix, output_path, **kwargs):
        df = self.predict_batch(active_matrix, **kwargs)
        path = Path(output_path)
        if path.suffix.lower() == ".xlsx":
            self._write_excel(df if isinstance(df, pd.DataFrame) else df.to_pandas(), path)
        elif path.suffix.lower() == ".csv":
            df.write_csv(str(path)) if hasattr(df, 'write_csv') else df.to_csv(path, index=False)
        return df

    def stream_export(self, active_matrix, output_path, chunk_size=5_000, **kwargs):
        path = Path(output_path)
        if path.suffix.lower() != ".csv": raise ValueError("stream_export only supports .csv")
        active_matrix = np.atleast_2d(np.asarray(active_matrix, dtype=int))
        self._validate_matrix(active_matrix)
        header = True
        for start in range(0, len(active_matrix), chunk_size):
            chunk_df = self.predict_batch(active_matrix[start:start+chunk_size], **kwargs)
            chunk_df.to_csv(path, mode='a' if not header else 'w', header=header, index=False)
            header = False

    def _rho_eff_for_active(self, active_indices: list[int]) -> float:
        if len(active_indices) < 2: return 0.0
        pairs = [self._corr_matrix[i, j] for k, i in enumerate(active_indices) for j in active_indices[k+1:]]
        return float(np.mean(pairs))

    def _blended_risk(self, active_indices: list[int], source_means: list[float]) -> float:
        if not active_indices: return 0.0
        mus = [source_means[i] for i in active_indices]
        if len(mus) == 1: return mus[0]
        rho = self._rho_eff_for_active(active_indices)
        return (1.0 - rho) * (1.0 - math.prod(1.0 - m for m in mus)) + rho * max(mus)

    def _blended_marginals(self, active_indices: list[int], source_means: list[float]) -> dict[int, float]:
        R_all = self._blended_risk(active_indices, source_means)
        contribs = {}
        for idx in active_indices:
            without = [i for i in active_indices if i != idx]
            R_without = self._blended_risk(without, source_means)
            contribs[idx] = round(R_all - R_without, 4)
        return contribs

    def _build_batch_pool(self, seed: Optional[int] = None) -> None:
        rng = np.random.default_rng(seed)
        self._batch_pool = _build_copula_pool(self._mu_values, self._kappa_values, self._corr_matrix, _BATCH_POOL_SIZE, rng)
        self._batch_source_means = self._batch_pool.mean(axis=1).astype(float)

    def _validate_matrix(self, matrix: np.ndarray):
        if matrix.ndim != 2 or matrix.shape[1] != self._n_sources or not np.isin(matrix, [0, 1]).all():
            raise ValueError(f"Invalid matrix shape or values. Expected (N, {self._n_sources}) with 0/1.")

    def _validate_flags(self, flags: np.ndarray):
        if flags.shape != (self._n_sources,) or not np.isin(flags, [0, 1]).all():
            raise ValueError(f"Invalid flag vector. Expected length {self._n_sources} with 0/1.")

    @staticmethod
    def _validate_source_params(mu: float, kappa: float, name: str):
        if not (0.0 < mu < 1.0) or kappa <= 0.0:
            raise ValueError(f"Invalid params for {name}: mu={mu}, kappa={kappa}")

    def _risk_band_label(self, score: float) -> Optional[str]:
        for threshold, label in self._risk_bands:
            if score >= threshold: return label
        return None

    def _format_single(self, sim: dict, flags: np.ndarray, label: Optional[str]) -> dict:
        contribs = sim["marginal_contribs"]
        active_idx = [i for i, a in enumerate(flags) if a]
        ranked = sorted(active_idx, key=lambda i: contribs.get(i, 0.0), reverse=True)
        others = [{"source": self._names[i], "marginal_contribution": round(contribs.get(i, 0.0), 4), "source_mean": sim["source_means"][i]} for i in ranked[1:]]
        return {
            "label": label, "risk_score": sim["mean_risk"], "median_risk": sim["median_risk"],
            "p5": sim["p5"], "p95": sim["p95"], "certainty": sim["certainty"],
            "certainty_label": sim["certainty_label"], "risk_band": self._risk_band_label(sim["mean_risk"]),
            "primary_driver": self._names[ranked[0]] if ranked else None,
            "primary_driver_contribution": round(contribs.get(ranked[0], 0.0), 4) if ranked else None,
            "other_matches": others, "source_means": sim["source_means"],
            "active_sources": [self._names[i] for i in active_idx],
        }

    def _format_row(self, sim: dict, flags: np.ndarray, label: str) -> dict:
        res = self._format_single(sim, flags, label)
        row = {
            "Label": res["label"], "Risk_Score": res["risk_score"], "Median_Risk": res["median_risk"],
            "P5": res["p5"], "P95": res["p95"], "Certainty": res["certainty"], "Certainty_Label": res["certainty_label"],
            "Primary_Driver": res["primary_driver"] or "", "Primary_Driver_Contribution": res["primary_driver_contribution"] or 0.0,
        }
        if self._risk_bands: row["Risk_Band"] = res["risk_band"] or ""
        for i, name in enumerate(self._names):
            col = name.replace(" ", "_")
            row[f"{col}_Active"] = int(flags[i])
            row[f"{col}_MarginalContrib"] = round(sim["marginal_contribs"].get(i, 0.0), 4) if flags[i] else 0.0
        row["Active_Sources"] = ", ".join(res["active_sources"]) if res["active_sources"] else "None"
        return row

    def _to_polars(self, rows: list[dict]) -> pl.DataFrame:
        if not rows or not _POLARS_AVAILABLE: return pl.DataFrame()
        str_cols = {"Label", "Certainty_Label", "Primary_Driver", "Active_Sources", "Risk_Band"}
        series = {}
        for col in rows[0]:
            vals = [r[col] for r in rows]
            dtype = pl.Utf8 if col in str_cols else (pl.Int32 if col.endswith("_Active") else pl.Float64)
            series[col] = pl.Series(col, vals, dtype=dtype)
        return pl.DataFrame(series)

    def print_result(self, result: dict):
        print(f"\n{'='*40}\nRisk Result | {result.get('label') or 'Record'}\n{'='*40}")
        print(f"Risk Score: {result['risk_score']:.1%} ({result['certainty_label']})")
        print(f"Driver: {result['primary_driver']} (+{result['primary_driver_contribution']:.1%})")
        print(f"Interval: [{result['p5']:.1%}, {result['p95']:.1%}]\n{'='*40}")

    def save_config(self, sources_dict: dict, output_path: Union[str, Path], version: Optional[str] = None, 
                    n_samples: Optional[int] = None, bump: str = "minor", correlations: Optional[list[dict]] = None):
        new_v = version or self._bump_version(self._version, bump)
        sources_list = []
        for name, p in sources_dict.items():
            self._validate_source_params(p["mu"], p["kappa"], name)
            sources_list.append({"name": name, "mu": float(p["mu"]), "kappa": float(p["kappa"])})
        
        cfg = {
            "version": new_v, "version_date": str(date.today()),
            "n_monte_carlo_samples": n_samples or self._n_samples_default,
            "sources": sources_list, "risk_bands": [{"label": l, "min_score": t} for t, l in self._risk_bands],
            "correlations": correlations if correlations is not None else self._corr_pairs
        }
        with open(output_path, "w") as f: json.dump(cfg, f, indent=2)

    @staticmethod
    def _bump_version(current: str, part: str) -> str:
        try: maj, min_, pat = (int(x) for x in current.split("."))
        except: return "1.0.0"
        if part == "major": return f"{maj+1}.0.0"
        if part == "minor": return f"{maj}.{min_+1}.0"
        return f"{maj}.{min_}.{pat+1}"

    def _write_excel(self, df: pd.DataFrame, path: Path):
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill, Font
        wb = Workbook(); ws = wb.active; ws.title = "Risk Results"
        hdr_f = Font(bold=True, color="FFFFFF"); hdr_b = PatternFill("solid", fgColor="1D6FA4")
        for c, col in enumerate(df.columns, 1):
            cell = ws.cell(1, c, col); cell.font = hdr_f; cell.fill = hdr_b
        for r, row in enumerate(df.values, 2):
            for c, val in enumerate(row, 1): ws.cell(r, c, val)
        wb.save(path)

    def __repr__(self) -> str:
        return f"NoisyORModelCopula(v{self._version}, sources={self._n_sources}, corr_pairs={len(self._corr_pairs)})"