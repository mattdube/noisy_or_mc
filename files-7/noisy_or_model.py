"""
noisy_or_model.py
=================
Noisy-OR Bayesian Risk Model — server / notebook edition.

Mirrors the logic of the Flask web app exactly, but designed for
programmatic use: single predictions, batch predictions, and
human-readable CSV / Excel exports.

Usage
-----
    from noisy_or_model import NoisyORModel

    model = NoisyORModel("noisy_or_config.json")

    # Single prediction (numpy array: 1 = source matched, 0 = did not match)
    import numpy as np
    result = model.predict(np.array([1, 0, 1, 0, 0, 1, 0, 0]))
    print(result)

    # Batch prediction
    inputs = np.array([
        [1, 0, 1, 0, 0, 1, 0, 0],
        [0, 1, 0, 1, 0, 0, 1, 0],
        [1, 1, 0, 0, 1, 0, 0, 1],
    ])
    df = model.predict_batch(inputs)
    print(df)

    # Export batch to Excel or CSV
    model.export_batch(inputs, "results.xlsx")
    model.export_batch(inputs, "results.csv")

Adding Sources
--------------
Edit noisy_or_config.json — add a dict to "sources":
    {"name": "My New Source", "mu": 0.25, "kappa": 6.0}
No code changes required.

Changing Monte Carlo sample count
----------------------------------
Edit "n_monte_carlo_samples" in noisy_or_config.json, or pass
n_samples= directly to predict() / predict_batch().
"""

from __future__ import annotations

import json
import math
import textwrap
from pathlib import Path
from typing import Optional, Union

import numpy as np
import pandas as pd


# ---------------------------------------------------------------------------
# Certainty thresholds — mirrors the web app's colour logic
# ---------------------------------------------------------------------------
def _certainty_label(cert: float) -> str:
    if cert >= 0.80:
        return "High"
    if cert >= 0.50:
        return "Medium"
    return "Low"


# ---------------------------------------------------------------------------
# Core simulation (identical math to the Flask run_simulation function)
# ---------------------------------------------------------------------------
def _run_simulation(
    mu_values: np.ndarray,
    kappa_values: np.ndarray,
    active_mask: np.ndarray,
    n_samples: int,
    rng: Optional[np.random.Generator] = None,
) -> dict:
    """
    Parameters
    ----------
    mu_values     : shape (n_sources,) — prior mean for each source
    kappa_values  : shape (n_sources,) — concentration for each source
    active_mask   : shape (n_sources,) — 1 if source matched, 0 otherwise
    n_samples     : number of Monte Carlo draws
    rng           : optional numpy RNG for reproducibility

    Returns
    -------
    dict with simulation statistics and per-source marginal contributions
    """
    if rng is None:
        rng = np.random.default_rng()

    n_sources = len(mu_values)

    # Sample Beta distributions for every source
    all_probs: list[np.ndarray] = []
    for mu, kappa in zip(mu_values, kappa_values):
        alpha = max(float(mu) * float(kappa), 0.001)
        beta  = max((1.0 - float(mu)) * float(kappa), 0.001)
        all_probs.append(rng.beta(alpha, beta, n_samples))

    # Combine only active sources via Noisy-OR
    active_probs = [p for p, a in zip(all_probs, active_mask) if a]

    if active_probs:
        combined = 1.0 - np.prod([1.0 - p for p in active_probs], axis=0)
    else:
        combined = np.zeros(n_samples)

    mean_risk   = float(np.mean(combined))
    median_risk = float(np.median(combined))
    p5          = float(np.percentile(combined, 5))
    p95         = float(np.percentile(combined, 95))
    certainty   = max(0.0, 1.0 - (p95 - p5))

    # Per-source marginal contributions (analytical from source means)
    source_means = [float(np.mean(p)) for p in all_probs]
    active_indices = [i for i, a in enumerate(active_mask) if a]

    contribs: dict[int, float] = {}
    if active_indices:
        active_mus = [source_means[i] for i in active_indices]
        prod_all = math.prod(1.0 - m for m in active_mus)
        overall_risk_approx = 1.0 - prod_all

        for idx, mu_i in zip(active_indices, active_mus):
            if mu_i > 0.9999:
                prod_without = 0.0
            else:
                prod_without = prod_all / (1.0 - mu_i)
            risk_without = 1.0 - prod_without
            contribs[idx] = overall_risk_approx - risk_without

    return {
        "mean_risk":      round(mean_risk, 4),
        "median_risk":    round(median_risk, 4),
        "p5":             round(p5, 4),
        "p95":            round(p95, 4),
        "certainty":      round(certainty, 4),
        "certainty_label": _certainty_label(certainty),
        "source_means":   [round(m, 4) for m in source_means],
        "marginal_contribs": contribs,   # {source_index: delta}
    }


# ---------------------------------------------------------------------------
# Public model class
# ---------------------------------------------------------------------------
class NoisyORModel:
    """
    Noisy-OR Bayesian Risk Model.

    Parameters
    ----------
    config_path : path to JSON config file (see noisy_or_config.json)
    """

    def __init__(self, config_path: Union[str, Path] = "noisy_or_config.json"):
        config_path = Path(config_path)
        if not config_path.exists():
            raise FileNotFoundError(f"Config not found: {config_path}")

        with open(config_path) as f:
            cfg = json.load(f)

        self._sources: list[dict] = cfg["sources"]
        self._n_samples_default: int = int(cfg.get("n_monte_carlo_samples", 20_000))

        self._names      = [s["name"]  for s in self._sources]
        self._mu_values  = np.array([s["mu"]    for s in self._sources], dtype=float)
        self._kappa_values = np.array([s["kappa"] for s in self._sources], dtype=float)
        self._n_sources  = len(self._sources)

    # ------------------------------------------------------------------
    # Single prediction
    # ------------------------------------------------------------------
    def predict(
        self,
        active_flags: np.ndarray,
        n_samples: Optional[int] = None,
        seed: Optional[int] = None,
        label: Optional[str] = None,
    ) -> dict:
        """
        Run the Noisy-OR model for a single input vector.

        Parameters
        ----------
        active_flags : 1-D numpy array of length n_sources.
                       1 = source matched this record, 0 = did not match.
        n_samples    : Monte Carlo draws (defaults to config value)
        seed         : optional int for reproducible results
        label        : optional string label for this record

        Returns
        -------
        dict with keys:
            label, risk_score, median_risk, p5, p95,
            certainty, certainty_label,
            primary_driver, primary_driver_contribution,
            other_matches  (list of dicts),
            source_means   (list of floats, one per source),
            active_sources (list of source names that matched)
        """
        active_flags = np.asarray(active_flags, dtype=int)
        self._validate_flags(active_flags)

        rng = np.random.default_rng(seed)
        n = n_samples or self._n_samples_default

        sim = _run_simulation(
            self._mu_values,
            self._kappa_values,
            active_flags,
            n,
            rng,
        )

        return self._format_single(sim, active_flags, label)

    # ------------------------------------------------------------------
    # Batch prediction
    # ------------------------------------------------------------------
    def predict_batch(
        self,
        active_matrix: np.ndarray,
        n_samples: Optional[int] = None,
        seed: Optional[int] = None,
        labels: Optional[list[str]] = None,
    ) -> pd.DataFrame:
        """
        Run the model for many input vectors.

        Parameters
        ----------
        active_matrix : 2-D numpy array, shape (n_records, n_sources).
                        Each row is one record; values are 0 or 1.
        n_samples     : Monte Carlo draws (defaults to config value)
        seed          : optional int; seeds are derived deterministically
                        per row so results are reproducible
        labels        : optional list of string labels for each row

        Returns
        -------
        pandas DataFrame with one row per record and columns:
            Label, Risk_Score, Median_Risk, P5, P95,
            Certainty, Certainty_Label,
            Primary_Driver, Primary_Driver_Contribution,
            <SourceName>_Active, <SourceName>_MarginalContrib  (per source),
            Active_Sources
        """
        active_matrix = np.atleast_2d(np.asarray(active_matrix, dtype=int))
        n_records = active_matrix.shape[0]
        n = n_samples or self._n_samples_default

        if labels is None:
            labels = [f"Record_{i+1}" for i in range(n_records)]

        rows = []
        for i, (row_flags, lbl) in enumerate(zip(active_matrix, labels)):
            self._validate_flags(row_flags)
            row_seed = None if seed is None else seed + i
            rng = np.random.default_rng(row_seed)
            sim = _run_simulation(
                self._mu_values, self._kappa_values, row_flags, n, rng
            )
            rows.append(self._format_row(sim, row_flags, lbl))

        return pd.DataFrame(rows)

    # ------------------------------------------------------------------
    # Export helpers
    # ------------------------------------------------------------------
    def export_batch(
        self,
        active_matrix: np.ndarray,
        output_path: Union[str, Path],
        n_samples: Optional[int] = None,
        seed: Optional[int] = None,
        labels: Optional[list[str]] = None,
    ) -> pd.DataFrame:
        """
        Run batch predictions and save to Excel (.xlsx) or CSV (.csv).

        Returns the DataFrame so you can inspect it in a notebook without
        loading the file again.
        """
        df = self.predict_batch(active_matrix, n_samples=n_samples,
                                seed=seed, labels=labels)
        output_path = Path(output_path)
        suffix = output_path.suffix.lower()

        if suffix == ".xlsx":
            self._write_excel(df, output_path)
        elif suffix == ".csv":
            df.to_csv(output_path, index=False)
            print(f"CSV saved → {output_path}")
        else:
            raise ValueError(f"Unsupported extension '{suffix}'. Use .xlsx or .csv")

        return df

    # ------------------------------------------------------------------
    # Internal helpers
    # ------------------------------------------------------------------
    def _validate_flags(self, flags: np.ndarray) -> None:
        if flags.shape != (self._n_sources,):
            raise ValueError(
                f"Expected array of length {self._n_sources} "
                f"(one entry per source), got shape {flags.shape}."
            )
        if not set(flags.tolist()).issubset({0, 1}):
            raise ValueError("active_flags must contain only 0 or 1.")

    def _format_single(self, sim: dict, flags: np.ndarray, label: Optional[str]) -> dict:
        contribs = sim["marginal_contribs"]
        active_idx = [i for i, a in enumerate(flags) if a]

        # Rank active sources by marginal contribution
        ranked = sorted(active_idx, key=lambda i: contribs.get(i, 0.0), reverse=True)

        primary_driver = None
        primary_contrib = None
        other_matches = []

        if ranked:
            primary_driver = self._names[ranked[0]]
            primary_contrib = round(contribs.get(ranked[0], 0.0), 4)
            for i in ranked[1:]:
                other_matches.append({
                    "source": self._names[i],
                    "marginal_contribution": round(contribs.get(i, 0.0), 4),
                    "source_mean": sim["source_means"][i],
                })

        return {
            "label":                      label,
            "risk_score":                 sim["mean_risk"],
            "median_risk":                sim["median_risk"],
            "p5":                         sim["p5"],
            "p95":                        sim["p95"],
            "certainty":                  sim["certainty"],
            "certainty_label":            sim["certainty_label"],
            "primary_driver":             primary_driver,
            "primary_driver_contribution": primary_contrib,
            "other_matches":              other_matches,
            "source_means":               sim["source_means"],
            "active_sources":             [self._names[i] for i in active_idx],
        }

    def _format_row(self, sim: dict, flags: np.ndarray, label: str) -> dict:
        """Flatten a single result into a dict of scalar/string columns."""
        res = self._format_single(sim, flags, label)
        contribs = sim["marginal_contribs"]

        row: dict = {
            "Label":                       res["label"],
            "Risk_Score":                  res["risk_score"],
            "Median_Risk":                 res["median_risk"],
            "P5":                          res["p5"],
            "P95":                         res["p95"],
            "Certainty":                   res["certainty"],
            "Certainty_Label":             res["certainty_label"],
            "Primary_Driver":              res["primary_driver"] or "",
            "Primary_Driver_Contribution": res["primary_driver_contribution"] or 0.0,
        }

        for i, name in enumerate(self._names):
            col_base = name.replace(" ", "_")
            row[f"{col_base}_Active"]           = int(flags[i])
            row[f"{col_base}_MarginalContrib"]  = round(contribs.get(i, 0.0), 4) if flags[i] else 0.0

        row["Active_Sources"] = ", ".join(res["active_sources"]) if res["active_sources"] else "None"
        return row

    def _write_excel(self, df: pd.DataFrame, path: Path) -> None:
        """Write a polished, human-readable Excel workbook."""
        from openpyxl import Workbook
        from openpyxl.styles import (
            Font, PatternFill, Alignment, Border, Side, numbers
        )
        from openpyxl.utils import get_column_letter

        wb = Workbook()

        # ---- Sheet 1: Results summary ----
        ws = wb.active
        ws.title = "Risk Results"

        # Colour palette
        HDR_FILL    = PatternFill("solid", fgColor="1D6FA4")
        SUBHDR_FILL = PatternFill("solid", fgColor="2E86C1")
        HIGH_FILL   = PatternFill("solid", fgColor="D5F5E3")
        MED_FILL    = PatternFill("solid", fgColor="FEF9E7")
        LOW_FILL    = PatternFill("solid", fgColor="FDEDEC")
        ALT_FILL    = PatternFill("solid", fgColor="EBF5FB")
        WHITE_FILL  = PatternFill("solid", fgColor="FFFFFF")

        HDR_FONT    = Font(name="Arial", bold=True, color="FFFFFF", size=10)
        BODY_FONT   = Font(name="Arial", size=10)
        BOLD_FONT   = Font(name="Arial", bold=True, size=10)
        TITLE_FONT  = Font(name="Arial", bold=True, size=13, color="1D6FA4")

        thin = Side(style="thin", color="BDBDBD")
        BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
        CENTER = Alignment(horizontal="center", vertical="center")
        LEFT   = Alignment(horizontal="left",   vertical="center")

        PCT_FMT   = "0.0%"
        SCORE_FMT = "0.0000"

        # Title row
        ws.merge_cells("A1:O1")
        title_cell = ws["A1"]
        title_cell.value     = "Noisy-OR Bayesian Risk Model — Batch Results"
        title_cell.font      = TITLE_FONT
        title_cell.alignment = LEFT
        ws.row_dimensions[1].height = 24

        # Column groups
        summary_cols = [
            ("Label",                       "Label",                   18, None),
            ("Risk_Score",                  "Risk Score",              11, PCT_FMT),
            ("Median_Risk",                 "Median Risk",             12, PCT_FMT),
            ("P5",                          "P5 (5th pct)",            11, PCT_FMT),
            ("P95",                         "P95 (95th pct)",          13, PCT_FMT),
            ("Certainty",                   "Certainty",               11, PCT_FMT),
            ("Certainty_Label",             "Certainty Level",         14, None),
            ("Primary_Driver",              "Primary Driver",          18, None),
            ("Primary_Driver_Contribution", "Primary Contribution",    18, PCT_FMT),
            ("Active_Sources",              "All Matched Sources",     30, None),
        ]

        source_cols = []
        for name in self._names:
            col_base = name.replace(" ", "_")
            source_cols.append((f"{col_base}_Active",          f"{name}\nActive",        10, None))
            source_cols.append((f"{col_base}_MarginalContrib", f"{name}\nMarginal Contrib", 14, PCT_FMT))

        all_cols = summary_cols + source_cols
        n_cols   = len(all_cols)

        # Header row (row 2)
        for col_num, (_, hdr, width, _fmt) in enumerate(all_cols, start=1):
            cell = ws.cell(row=2, column=col_num, value=hdr)
            cell.font      = HDR_FONT
            cell.fill      = HDR_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border    = BORDER
            ws.column_dimensions[get_column_letter(col_num)].width = width
        ws.row_dimensions[2].height = 36

        # Data rows
        cert_fill_map = {"High": HIGH_FILL, "Medium": MED_FILL, "Low": LOW_FILL}

        for row_num, (_, row_data) in enumerate(df.iterrows(), start=3):
            cert_lbl  = str(row_data.get("Certainty_Label", ""))
            row_fill  = cert_fill_map.get(cert_lbl, WHITE_FILL)
            alt_fill  = ALT_FILL if (row_num % 2 == 0) else WHITE_FILL

            for col_num, (df_col, _, _, fmt) in enumerate(all_cols, start=1):
                val  = row_data.get(df_col, "")
                cell = ws.cell(row=row_num, column=col_num, value=val)
                cell.font      = BODY_FONT
                cell.border    = BORDER
                cell.alignment = CENTER if fmt else LEFT

                # Shade certainty columns
                if df_col in ("Certainty", "Certainty_Label"):
                    cell.fill = row_fill
                    if df_col == "Certainty_Label":
                        cell.font = Font(name="Arial", bold=True, size=10)
                else:
                    cell.fill = alt_fill

                if fmt:
                    cell.number_format = fmt

        ws.freeze_panes = "A3"
        ws.auto_filter.ref = f"A2:{get_column_letter(n_cols)}2"

        # ---- Sheet 2: Config snapshot ----
        ws2 = wb.create_sheet("Model Config")
        ws2["A1"].value = "Noisy-OR Model Configuration"
        ws2["A1"].font  = TITLE_FONT

        cfg_hdrs = ["Source Name", "Mu (Prior Mean)", "Kappa (Concentration)"]
        for c, h in enumerate(cfg_hdrs, start=1):
            cell = ws2.cell(row=2, column=c, value=h)
            cell.font = HDR_FONT
            cell.fill = HDR_FILL
            cell.border = BORDER
            cell.alignment = CENTER

        ws2.column_dimensions["A"].width = 20
        ws2.column_dimensions["B"].width = 16
        ws2.column_dimensions["C"].width = 20

        for r, src in enumerate(self._sources, start=3):
            ws2.cell(row=r, column=1, value=src["name"]).font   = BODY_FONT
            ws2.cell(row=r, column=2, value=src["mu"]).font     = BODY_FONT
            ws2.cell(row=r, column=3, value=src["kappa"]).font  = BODY_FONT
            for c in range(1, 4):
                ws2.cell(row=r, column=c).border = BORDER

        n_row = len(self._sources) + 4
        ws2.cell(row=n_row, column=1, value="Monte Carlo Samples").font = BOLD_FONT
        ws2.cell(row=n_row, column=2, value=self._n_samples_default).font = BODY_FONT

        wb.save(path)
        print(f"Excel workbook saved → {path}")

    # ------------------------------------------------------------------
    # Pretty-print for single prediction (notebook / script use)
    # ------------------------------------------------------------------
    def print_result(self, result: dict) -> None:
        """Pretty-print a single prediction result."""
        lbl = result.get("label") or "Record"
        print(f"\n{'='*60}")
        print(f"  NOISY-OR RISK RESULT  |  {lbl}")
        print(f"{'='*60}")
        print(f"  Risk Score   : {result['risk_score']:.1%}")
        print(f"  Median Risk  : {result['median_risk']:.1%}")
        print(f"  90% Interval : [{result['p5']:.1%},  {result['p95']:.1%}]")
        print(f"  Certainty    : {result['certainty']:.1%}  ({result['certainty_label']})")
        print(f"\n  Primary Driver : {result['primary_driver']}")
        print(f"  Primary Contribution : +{result['primary_driver_contribution']:.1%}")

        if result["other_matches"]:
            print(f"\n  Other Matched Sources:")
            for m in result["other_matches"]:
                print(f"    • {m['source']:<20}  marginal +{m['marginal_contribution']:.1%}")
        else:
            print("\n  No other matched sources.")

        print(f"\n  Active sources : {', '.join(result['active_sources']) or 'None'}")
        print(f"{'='*60}\n")

    @property
    def source_names(self) -> list[str]:
        return self._names

    @property
    def n_sources(self) -> int:
        return self._n_sources

    def __repr__(self) -> str:
        return (
            f"NoisyORModel(n_sources={self._n_sources}, "
            f"n_samples={self._n_samples_default})"
        )
